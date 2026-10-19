import os
import re
import docx
import docx.opc.constants
from docx import Document
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_border(cell, **kwargs):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(r'''
        <w:tcBorders {} >
            <w:top w:val="{top}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:left w:val="{left}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:bottom w:val="{bottom}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:right w:val="{right}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        </w:tcBorders>
    '''.format(nsdecls('w'),
               top=kwargs.get('top', 'single'),
               bottom=kwargs.get('bottom', 'single'),
               left=kwargs.get('left', 'single'),
               right=kwargs.get('right', 'single'),
               sz=kwargs.get('sz', '4'),
               color=kwargs.get('color', 'D3D3D3')))
    tcPr.append(tcBorders)

def set_cell_shading(cell, color_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    tcPr.append(shd)

def clean_math_text(text):
    if not text:
        return text
    # Ký hiệu LaTeX phổ biến
    text = text.replace(r'\times', '×')
    text = text.replace(r'\le', '≤')
    text = text.replace(r'\ge', '≥')
    text = text.replace(r'^\circ', '°')
    text = text.replace(r'\Delta', 'Δ')
    text = text.replace(r'\rightarrow', '→')
    text = text.replace(r'\leftarrow', '←')
    text = text.replace(r'\approx', '≈')
    text = text.replace(r'\pm', '±')
    text = text.replace(r'\cdot', '·')
    # Bỏ dấu $ bao bọc: $X$ -> X
    text = re.sub(r'\$([^$]+)\$', r'\1', text)
    return text

def add_omml_equation(doc, math_expr):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(8)
    
    def create_sub(var, sub):
        return f'<m:sSub><m:e><m:r><m:rPr><m:sty m:val="i"/></m:rPr><m:t>{var}</m:t></m:r></m:e><m:sub><m:r><m:t>{sub}</m:t></m:r></m:sub></m:sSub>'
        
    if 'EAR' in math_expr:
        ear_num = f'<m:r><m:t>|</m:t></m:r>{create_sub("p", "159")}<m:r><m:t> - </m:t></m:r>{create_sub("p", "145")}<m:r><m:t>| + |</m:t></m:r>{create_sub("p", "158")}<m:r><m:t> - </m:t></m:r>{create_sub("p", "153")}<m:r><m:t>|</m:t></m:r>'
        ear_den = f'<m:r><m:t>2 × |</m:t></m:r>{create_sub("p", "33")}<m:r><m:t> - </m:t></m:r>{create_sub("p", "133")}<m:r><m:t>|</m:t></m:r>'
        omml = f'''<m:oMathPara xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">
          <m:oMath>
            <m:r><m:rPr><m:sty m:val="bi"/></m:rPr><m:t>EAR</m:t></m:r>
            <m:r><m:t> = </m:t></m:r>
            <m:f>
              <m:num>{ear_num}</m:num>
              <m:den>{ear_den}</m:den>
            </m:f>
          </m:oMath>
        </m:oMathPara>'''
        p._p.append(parse_xml(omml))
    elif 'MAR' in math_expr:
        mar_num = f'<m:r><m:t>|</m:t></m:r>{create_sub("p", "13")}<m:r><m:t> - </m:t></m:r>{create_sub("p", "14")}<m:r><m:t>|</m:t></m:r>'
        mar_den = f'<m:r><m:t>|</m:t></m:r>{create_sub("p", "61")}<m:r><m:t> - </m:t></m:r>{create_sub("p", "291")}<m:r><m:t>|</m:t></m:r>'
        omml = f'''<m:oMathPara xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">
          <m:oMath>
            <m:r><m:rPr><m:sty m:val="bi"/></m:rPr><m:t>MAR</m:t></m:r>
            <m:r><m:t> = </m:t></m:r>
            <m:f>
              <m:num>{mar_num}</m:num>
              <m:den>{mar_den}</m:den>
            </m:f>
          </m:oMath>
        </m:oMathPara>'''
        p._p.append(parse_xml(omml))
    else:
        cleaned = clean_math_text(math_expr)
        run = p.add_run(cleaned)
        run.font.name = 'Cambria Math'
        run.font.size = Pt(13)
        run.font.italic = True

CITATION_MAP = {
    '[3]': 'https://thuvienphapluat.vn/van-ban/Lao-dong-Tien-luong/Bo-Luat-lao-dong-2019-333670.aspx',
    '[4]': 'https://thuvienphapluat.vn/van-ban/Cong-nghe-thong-tin/Nghi-dinh-13-2023-ND-CP-bao-ve-du-lieu-ca-nhan-536968.aspx',
    '[5]': 'https://thuvienphapluat.vn/van-ban/Lao-dong-Tien-luong/Nghi-dinh-145-2020-ND-CP-huong-dan-Bo-luat-Lao-dong-ve-dieu-kien-quan-he-lao-dong-460987.aspx',
    '[6]': 'https://thuvienphapluat.vn/',
    '[8]': 'https://ai.google.dev/edge/mediapipe/solutions/vision/face_landmarker',
    '[9]': 'https://ai.google.dev/gemini-api/docs/function-calling',
    '[10]': 'https://scikit-learn.org/stable/',
    '[11]': 'https://www.kaggle.com/datasets/pavansubhasht/ibm-hr-analytics-attrition-dataset',
    '[12]': 'https://www.mongodb.com/docs/manual/core/document/',
    '[13]': 'https://nodejs.org/en/learn/asynchronous-work/event-loop-timers-and-nexttick',
    '[14]': 'https://react.dev/learn/preserving-and-resetting-state',
}

def add_hyperlink(paragraph, url, text, font_size=Pt(13), is_bold=False, is_italic=False, color="0563C1"):
    try:
        part = paragraph.part
        r_id = part.relate_to(url, docx.opc.constants.RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
        hyperlink = parse_xml(f'<w:hyperlink {nsdecls("w")} r:id="{r_id}" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"/>')
        new_run = parse_xml(f'<w:r {nsdecls("w")}/>')
        rPr = parse_xml(f'<w:rPr {nsdecls("w")}/>')
        
        rFonts = parse_xml(f'<w:rFonts {nsdecls("w")} w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>')
        rPr.append(rFonts)
        
        sz_val = str(int(font_size.pt * 2))
        sz = parse_xml(f'<w:sz {nsdecls("w")} w:val="{sz_val}"/>')
        rPr.append(sz)
        
        c = parse_xml(f'<w:color {nsdecls("w")} w:val="{color}"/>')
        rPr.append(c)
        
        u = parse_xml(f'<w:u {nsdecls("w")} w:val="single"/>')
        rPr.append(u)
        
        if is_bold:
            rPr.append(parse_xml(f'<w:b {nsdecls("w")}/>'))
        if is_italic:
            rPr.append(parse_xml(f'<w:i {nsdecls("w")}/>'))
            
        rPr.append(parse_xml(f'<w:noProof {nsdecls("w")}/>'))
        new_run.append(rPr)
        t = parse_xml(f'<w:t {nsdecls("w")}>{text}</w:t>')
        new_run.append(t)
        hyperlink.append(new_run)
        paragraph._p.append(hyperlink)
        return hyperlink
    except Exception:
        r = paragraph.add_run(text)
        r.font.name = 'Times New Roman'
        r.font.size = font_size
        return r

def render_inline_markdown(paragraph, text, font_size=Pt(13), default_color=None, is_caption=False):
    if not text:
        return
    pattern = r'(\[.*?\]\(https?://[^\s)]+\)|\[\d+\]|https?://[^\s,;\)\]]+|\*\*.*?\*\*|\*.*?\*)'
    tokens = re.split(pattern, text)
    
    for token in tokens:
        if not token:
            continue
            
        link_m = re.match(r'^\[(.*?)\]\((https?://[^\s)]+)\)$', token)
        if link_m:
            anchor = link_m.group(1)
            url = link_m.group(2)
            add_hyperlink(paragraph, url, anchor, font_size=font_size, is_italic=is_caption)
            continue
            
        if token in CITATION_MAP:
            url = CITATION_MAP[token]
            add_hyperlink(paragraph, url, token, font_size=font_size, is_italic=is_caption)
            continue
            
        if token.startswith('http://') or token.startswith('https://'):
            clean_url = token.rstrip('.,;:)')
            trailing = token[len(clean_url):]
            add_hyperlink(paragraph, clean_url, clean_url, font_size=font_size, is_italic=is_caption)
            if trailing:
                r = paragraph.add_run(trailing)
                r.font.name = 'Times New Roman'
                r.font.size = font_size
                if default_color:
                    r.font.color.rgb = default_color
            continue
            
        if token.startswith('**') and token.endswith('**') and len(token) >= 4:
            r = paragraph.add_run(token[2:-2])
            r.font.name = 'Times New Roman'
            r.font.size = font_size
            r.font.bold = True
            if default_color:
                r.font.color.rgb = default_color
            continue
            
        if token.startswith('*') and token.endswith('*') and len(token) >= 2:
            r = paragraph.add_run(token[1:-1])
            r.font.name = 'Times New Roman'
            r.font.size = font_size
            r.font.italic = True
            if default_color:
                r.font.color.rgb = default_color
            continue
            
        r = paragraph.add_run(token)
        r.font.name = 'Times New Roman'
        r.font.size = font_size
        if is_caption:
            r.font.italic = True
        if default_color:
            r.font.color.rgb = default_color

def parse_markdown_to_doc(doc, md_content, base_dir=r'c:\K23CNT1\DANT\DATN\docs\report'):
    lines = md_content.split('\n')
    in_code_block = False
    code_block_lines = []
    in_table = False
    table_lines = []
    
    def flush_table(tbl_lines):
        if not tbl_lines:
            return
        rows_data = []
        for line in tbl_lines:
            if re.match(r'^\s*\|?\s*[-:]+[-| :]*$', line):
                continue
            cells = [c.strip() for c in line.strip().strip('|').split('|')]
            if cells:
                rows_data.append(cells)
        if not rows_data:
            return
            
        num_cols = max(len(r) for r in rows_data)
        table = doc.add_table(rows=len(rows_data), cols=num_cols)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False

        for i, row in enumerate(rows_data):
            for j, cell_text in enumerate(row):
                cell = table.cell(i, j)
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(4)
                p.paragraph_format.line_spacing = 1.15
                
                if i == 0:
                    set_cell_shading(cell, "1F497D")
                    set_cell_border(cell, color="1F497D", sz="8")
                    run = p.add_run(cell_text)
                    run.font.name = 'Times New Roman'
                    run.font.size = Pt(11)
                    run.font.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                else:
                    bg_color = "F2F5F9" if i % 2 == 1 else "FFFFFF"
                    set_cell_shading(cell, bg_color)
                    set_cell_border(cell, color="D9D9D9", sz="4")
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if j == 0 else WD_ALIGN_PARAGRAPH.LEFT
                    clean_text = cell_text
                    run = p.add_run(clean_text.replace('**', ''))
                    run.font.name = 'Times New Roman'
                    run.font.size = Pt(11)
                    if '**' in cell_text:
                        run.font.bold = True
                    run.font.color.rgb = RGBColor(30, 30, 30)

        p_space = doc.add_paragraph()
        p_space.paragraph_format.space_before = Pt(0)
        p_space.paragraph_format.space_after = Pt(4)

    def flush_code_block(code_lines):
        if not code_lines:
            return
        box_text = '\n'.join(code_lines)
        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.cell(0, 0)
        set_cell_shading(cell, "F8F9FA")
        set_cell_border(cell, color="4B6B94", sz="6")
        
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.05
        run = p.add_run(box_text)
        run.font.name = 'Consolas'
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(30, 41, 59)
        
        p_space = doc.add_paragraph()
        p_space.paragraph_format.space_before = Pt(0)
        p_space.paragraph_format.space_after = Pt(4)

    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        
        if stripped.startswith('```'):
            if in_code_block:
                flush_code_block(code_block_lines)
                code_block_lines = []
                in_code_block = False
            else:
                if in_table:
                    flush_table(table_lines)
                    table_lines = []
                    in_table = False
                in_code_block = True
            i += 1
            continue
            
        if in_code_block:
            code_block_lines.append(line)
            i += 1
            continue

        if stripped.startswith('|') and stripped.endswith('|'):
            in_table = True
            table_lines.append(stripped)
            i += 1
            continue
        else:
            if in_table:
                flush_table(table_lines)
                table_lines = []
                in_table = False

        if not stripped or stripped == '---':
            i += 1
            continue

        if stripped.startswith('# '):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(18)
            p.paragraph_format.space_after = Pt(14)
            run = p.add_run(stripped[2:])
            run.font.name = 'Times New Roman'
            run.font.size = Pt(16)
            run.font.bold = True
            run.font.color.rgb = RGBColor(16, 44, 87)
            i += 1
            continue

        if stripped.startswith('## '):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(stripped[3:])
            run.font.name = 'Times New Roman'
            run.font.size = Pt(14)
            run.font.bold = True
            run.font.color.rgb = RGBColor(16, 44, 87)
            i += 1
            continue

        if stripped.startswith('### '):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(stripped[4:])
            run.font.name = 'Times New Roman'
            run.font.size = Pt(13)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0, 0, 0)
            i += 1
            continue

        if stripped.startswith('#### '):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(stripped[5:])
            run.font.name = 'Times New Roman'
            run.font.size = Pt(13)
            run.font.bold = True
            run.font.italic = True
            run.font.color.rgb = RGBColor(40, 40, 40)
            i += 1
            continue

        if stripped.startswith('Bảng '):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(stripped)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(12)
            run.font.bold = True
            run.font.color.rgb = RGBColor(16, 44, 87)
            i += 1
            continue

        img_match = re.match(r'^!\[(.*?)\]\((.*?)\)', stripped)
        if img_match:
            caption = img_match.group(1)
            img_rel_path = img_match.group(2)
            img_path = os.path.normpath(os.path.join(base_dir, img_rel_path))
            if os.path.exists(img_path):
                p_img = doc.add_paragraph()
                p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_img.paragraph_format.space_before = Pt(8)
                p_img.paragraph_format.space_after = Pt(2)
                run_img = p_img.add_run()
                run_img.add_picture(img_path, width=Inches(5.8))
                
                if caption:
                    p_cap = doc.add_paragraph()
                    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p_cap.paragraph_format.space_before = Pt(2)
                    p_cap.paragraph_format.space_after = Pt(10)
                    render_inline_markdown(p_cap, caption, font_size=Pt(11), default_color=RGBColor(60, 60, 60), is_caption=True)
            i += 1
            continue

        if stripped.startswith('*Hình ') and stripped.endswith('*'):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(12)
            render_inline_markdown(p, stripped.strip('*'), font_size=Pt(12), default_color=RGBColor(60, 60, 60), is_caption=True)
            i += 1
            continue

        if stripped.startswith('$$') and stripped.endswith('$$') and len(stripped) > 4:
            math_expr = stripped[2:-2].strip()
            add_omml_equation(doc, math_expr)
            i += 1
            continue

        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.2
        
        if stripped.startswith('- ') or stripped.startswith('* '):
            p.paragraph_format.left_indent = Inches(0.25)
            text_to_parse = clean_math_text(stripped[2:])
            r_b = p.add_run('• ')
            r_b.font.name = 'Times New Roman'
            r_b.font.size = Pt(13)
            r_b.font.bold = True
        elif re.match(r'^\d+\.\s', stripped):
            match = re.match(r'^(\d+\.\s)(.*)', stripped)
            prefix = match.group(1)
            text_to_parse = clean_math_text(match.group(2))
            p.paragraph_format.left_indent = Inches(0.25)
            r_b = p.add_run(prefix)
            r_b.font.name = 'Times New Roman'
            r_b.font.size = Pt(13)
            r_b.font.bold = True
        elif stripped.startswith('+ '):
            p.paragraph_format.left_indent = Inches(0.4)
            text_to_parse = clean_math_text(stripped[2:])
            r_b = p.add_run('- ')
            r_b.font.name = 'Times New Roman'
            r_b.font.size = Pt(13)
        else:
            text_to_parse = clean_math_text(stripped)

        render_inline_markdown(p, text_to_parse, font_size=Pt(13))
        i += 1

    if in_table:
        flush_table(table_lines)
    if in_code_block:
        flush_code_block(code_block_lines)

def setup_page(doc):
    for section in doc.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(3.0)
        section.right_margin = Cm(2.0)
        
        footer = section.footer
        p_ft = footer.paragraphs[0]
        p_ft.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fldSimple = parse_xml(r'<w:fldSimple %s w:instr="PAGE"/>' % nsdecls('w'))
        p_ft._p.append(fldSimple)
        run_ft = p_ft.runs[0] if p_ft.runs else p_ft.add_run()
        run_ft.font.name = 'Times New Roman'
        run_ft.font.size = Pt(11)

    # 1. Tắt hoàn toàn gạch đỏ kiểm tra chính tả/ngữ pháp tiếng Anh trong Microsoft Word
    settings = doc.settings.element
    if settings.find(qn('w:hideSpellingErrors')) is None:
        settings.append(parse_xml(r'<w:hideSpellingErrors %s/>' % nsdecls('w')))
    if settings.find(qn('w:hideGrammaticalErrors')) is None:
        settings.append(parse_xml(r'<w:hideGrammaticalErrors %s/>' % nsdecls('w')))

    # 2. Cấu hình Normal style với tiếng Việt và cờ noProof (No Proofing)
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(13)
    normal_style.font.color.rgb = RGBColor(0, 0, 0)
    normal_style.paragraph_format.line_spacing = 1.2
    normal_style.paragraph_format.space_before = Pt(6)
    normal_style.paragraph_format.space_after = Pt(6)

    rPr = normal_style.element.get_or_add_rPr()
    lang = parse_xml(r'<w:lang %s w:val="vi-VN" w:eastAsia="vi-VN" w:bidi="vi-VN"/>' % nsdecls('w'))
    rPr.append(lang)
    noProof = parse_xml(r'<w:noProof %s/>' % nsdecls('w'))
    rPr.append(noProof)

def add_cover_page(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(36)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run("TRƯỜNG ĐẠI HỌC NGUYỄN TRÃI\nKHOA CÔNG NGHỆ THÔNG TIN & CHUYỂN ĐỔI SỐ")
    run.font.name = 'Times New Roman'
    run.font.size = Pt(14)
    run.font.bold = True
    run.font.color.rgb = RGBColor(16, 44, 87)

    p_line = doc.add_paragraph()
    p_line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_line.paragraph_format.space_before = Pt(0)
    p_line.paragraph_format.space_after = Pt(36)
    r_line = p_line.add_run("-----------------***-----------------")
    r_line.font.name = 'Times New Roman'
    r_line.font.size = Pt(12)
    r_line.font.bold = True

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(30)
    p_title.paragraph_format.space_after = Pt(10)
    r_t1 = p_title.add_run("ĐỒ ÁN TỐT NGHIỆP ĐẠI HỌC\n\n")
    r_t1.font.name = 'Times New Roman'
    r_t1.font.size = Pt(16)
    r_t1.font.bold = True
    r_t1.font.color.rgb = RGBColor(16, 44, 87)
    
    r_t2 = p_title.add_run("ĐỀ TÀI:\nPHÁT TRIỂN HỆ THỐNG QUẢN LÝ NHÂN SỰ DOANH NGHIỆP TÍCH HỢP TRÍ TUỆ NHÂN TẠO PEOPLEINSIGHT AI")
    r_t2.font.name = 'Times New Roman'
    r_t2.font.size = Pt(16)
    r_t2.font.bold = True
    r_t2.font.color.rgb = RGBColor(192, 0, 0)

    p_info = doc.add_paragraph()
    p_info.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_info.paragraph_format.left_indent = Inches(1.0)
    p_info.paragraph_format.space_before = Pt(60)
    p_info.paragraph_format.space_after = Pt(60)
    p_info.paragraph_format.line_spacing = 1.3
    
    info_text = (
        "Ngành:                      Công nghệ thông tin\n\n"
        "Giảng viên hướng dẫn:       ThS. Lê Trọng Nhất\n"
        "Sinh viên thực hiện:        Nguyễn Anh Tuấn\n"
        "Lớp:                        K23CNT1\n"
        "Mã sinh viên:               2310900113"
    )
    r_info = p_info.add_run(info_text)
    r_info.font.name = 'Times New Roman'
    r_info.font.size = Pt(13)
    r_info.font.bold = True

    p_date = doc.add_paragraph()
    p_date.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_date.paragraph_format.space_before = Pt(40)
    r_date = p_date.add_run("Hà Nội, 12/2026")
    r_date.font.name = 'Times New Roman'
    r_date.font.size = Pt(13)
    r_date.font.italic = True
    
    doc.add_page_break()

def add_sub_cover_page(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(36)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run("TRƯỜNG ĐẠI HỌC NGUYỄN TRÃI\nKHOA CÔNG NGHỆ THÔNG TIN & CHUYỂN ĐỔI SỐ")
    run.font.name = 'Times New Roman'
    run.font.size = Pt(14)
    run.font.bold = True
    run.font.color.rgb = RGBColor(16, 44, 87)

    p_line = doc.add_paragraph()
    p_line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_line.paragraph_format.space_before = Pt(0)
    p_line.paragraph_format.space_after = Pt(36)
    r_line = p_line.add_run("-----------------***-----------------")
    r_line.font.name = 'Times New Roman'
    r_line.font.size = Pt(12)
    r_line.font.bold = True

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(30)
    p_title.paragraph_format.space_after = Pt(10)
    r_t1 = p_title.add_run("ĐỒ ÁN TỐT NGHIỆP ĐẠI HỌC\n\n")
    r_t1.font.name = 'Times New Roman'
    r_t1.font.size = Pt(16)
    r_t1.font.bold = True
    r_t1.font.color.rgb = RGBColor(16, 44, 87)
    
    r_t2 = p_title.add_run("ĐỀ TÀI:\nPHÁT TRIỂN HỆ THỐNG QUẢN LÝ NHÂN SỰ DOANH NGHIỆP TÍCH HỢP TRÍ TUỆ NHÂN TẠO PEOPLEINSIGHT AI")
    r_t2.font.name = 'Times New Roman'
    r_t2.font.size = Pt(16)
    r_t2.font.bold = True
    r_t2.font.color.rgb = RGBColor(192, 0, 0)

    p_info = doc.add_paragraph()
    p_info.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_info.paragraph_format.left_indent = Inches(1.0)
    p_info.paragraph_format.space_before = Pt(40)
    p_info.paragraph_format.space_after = Pt(40)
    p_info.paragraph_format.line_spacing = 1.3
    
    info_text = (
        "Ngành:                      Công nghệ thông tin\n\n"
        "Giảng viên hướng dẫn:       ThS. Lê Trọng Nhất\n"
        "Chữ ký GVHD:                ........................................\n\n"
        "Sinh viên thực hiện:        Nguyễn Anh Tuấn\n"
        "Lớp:                        K23CNT1\n"
        "Mã sinh viên:               2310900113"
    )
    r_info = p_info.add_run(info_text)
    r_info.font.name = 'Times New Roman'
    r_info.font.size = Pt(13)
    r_info.font.bold = True

    p_date = doc.add_paragraph()
    p_date.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_date.paragraph_format.space_before = Pt(30)
    r_date = p_date.add_run("Hà Nội, 12/2026")
    r_date.font.name = 'Times New Roman'
    r_date.font.size = Pt(13)
    r_date.font.italic = True
    
    doc.add_page_break()

def main():
    output_dir = r'c:\K23CNT1\DANT\DATN\docs\report'
    
    # 1. File chỉ riêng Chương 1 kèm Tài liệu tham khảo
    doc_ch1 = Document()
    setup_page(doc_ch1)
    with open(os.path.join(output_dir, 'ch1_tong_quan.md'), 'r', encoding='utf-8') as f:
        parse_markdown_to_doc(doc_ch1, f.read())
        
    tai_lieu_path = os.path.join(output_dir, 'tai_lieu_tham_khao.md')
    if os.path.exists(tai_lieu_path):
        doc_ch1.add_page_break()
        with open(tai_lieu_path, 'r', encoding='utf-8') as f:
            parse_markdown_to_doc(doc_ch1, f.read())
            
    file_ch1 = os.path.join(output_dir, 'CHƯƠNG_1_TỔNG_QUAN_PEOPLEINSIGHT_AI.docx')
    doc_ch1.save(file_ch1)
    print(f"Created: {file_ch1}")

    # 2. File Đầy đủ: Bìa chính + Bìa phụ + Lời cảm ơn + Tóm tắt + Mở đầu + Chương 1 + Tài liệu tham khảo
    doc_full = Document()
    setup_page(doc_full)
    add_cover_page(doc_full)
    add_sub_cover_page(doc_full)
    
    # Lời cảm ơn
    loi_cam_on_path = os.path.join(output_dir, 'loi_cam_on.md')
    if os.path.exists(loi_cam_on_path):
        with open(loi_cam_on_path, 'r', encoding='utf-8') as f:
            parse_markdown_to_doc(doc_full, f.read())
        doc_full.add_page_break()

    # Tóm tắt nội dung đồ án
    tom_tat_path = os.path.join(output_dir, 'tom_tat.md')
    if os.path.exists(tom_tat_path):
        with open(tom_tat_path, 'r', encoding='utf-8') as f:
            parse_markdown_to_doc(doc_full, f.read())
        doc_full.add_page_break()
    
    # Phần Mở đầu
    with open(os.path.join(output_dir, 'mo_dau.md'), 'r', encoding='utf-8') as f:
        parse_markdown_to_doc(doc_full, f.read())
        
    doc_full.add_page_break()
    
    # Chương 1
    with open(os.path.join(output_dir, 'ch1_tong_quan.md'), 'r', encoding='utf-8') as f:
        parse_markdown_to_doc(doc_full, f.read())
        
    # Tài liệu tham khảo
    if os.path.exists(tai_lieu_path):
        doc_full.add_page_break()
        with open(tai_lieu_path, 'r', encoding='utf-8') as f:
            parse_markdown_to_doc(doc_full, f.read())
        
    file_full = os.path.join(output_dir, 'BAO_CAO_DATN_NTU_MO_DAU_VA_CHUONG_1.docx')
    doc_full.save(file_full)
    print(f"Created: {file_full}")

if __name__ == '__main__':
    main()
